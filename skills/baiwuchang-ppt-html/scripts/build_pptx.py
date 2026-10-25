# -*- coding: utf-8 -*-
"""
百无常 PPT —— 生成「可编辑」的 PowerPoint(.pptx)。
文字 = 原生文本框（可改）；封面主视觉 = 图片（可右键“更改图片”）；左下角 = 灰色 logo 图片。
颜色 / 字体 / 版式遵循《设计规范.md》：白底、黑色超粗标题、细正文、强调色点缀。
运行：python build_pptx.py    输出：../交付示例/baiwuchang-ppt.pptx
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsdecls
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
LOGO = os.path.join(ROOT, "html参考", "极简", "灰色logo2.png")
OUTDIR = os.path.join(ROOT, "交付示例"); os.makedirs(OUTDIR, exist_ok=True)
OUT = os.path.join(OUTDIR, "baiwuchang-ppt.pptx")
PLACE = os.path.join(HERE, "_image_placeholder.png")

DW, DH = 1672, 941
GREEN = RGBColor(0x39, 0xff, 0x14)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xff, 0xff, 0xff)
MUTED = RGBColor(0x6e, 0x6e, 0x6e)
DARK  = RGBColor(0x40, 0x40, 0x40)
LINE  = RGBColor(0xe2, 0xe2, 0xe2)
SOFTG = RGBColor(0xf2, 0xf2, 0xf2)
FONT  = "思源黑体"

def EX(px): return Emu(int(round(px * 12192000 / DW)))
def EY(px): return Emu(int(round(px * 6858000 / DH)))

def make_placeholder():
    w = h = 900
    img = Image.new("RGB", (w, h), (238, 238, 238))
    d = ImageDraw.Draw(img)
    d.rectangle([8, 8, w - 8, h - 8], outline=(199, 199, 199), width=3)
    try:
        f1 = ImageFont.truetype("msyh.ttc", 50); f2 = ImageFont.truetype("msyh.ttc", 30)
    except Exception:
        f1 = ImageFont.load_default(); f2 = f1
    d.text((w/2, h/2-36), "图片占位", fill=(120,120,120), anchor="mm", font=f1)
    d.text((w/2, h/2+34), "右键 → 更改图片", fill=(150,150,150), anchor="mm", font=f2)
    img.save(PLACE)

make_placeholder()
prs = Presentation()
prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]

def setfont(run, name=FONT):
    run.font.name = name
    run._r.get_or_add_rPr().append(parse_xml('<a:ea %s typeface="%s"/>' % (nsdecls('a'), name)))

def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.0, wrap=True):
    tb = slide.shapes.add_textbox(EX(x), EY(y), EX(w), EY(h)); tf = tb.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line != 1.0: p.line_spacing = line
        for (txt, size, bold, color) in para:
            r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; setfont(r)
    return tb

def rect(slide, x, y, w, h, fill=None, lc=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, soft=0):
    sp = slide.shapes.add_shape(shape, EX(x), EY(y), EX(w), EY(h)); sp.shadow.inherit = False
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if lc is None: sp.line.fill.background()
    else: sp.line.color.rgb = lc; sp.line.width = Pt(lw)
    if soft > 0:
        sp._element.spPr.append(parse_xml('<a:effectLst %s><a:softEdge rad="%d"/></a:effectLst>' % (nsdecls('a'), soft)))
    return sp

def gblob(slide, x, y, w, h, c1, c2):
    """柔和无边界渐变椭圆（软边羽化）——近似 HTML 的渐变光团。"""
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, EX(x), EY(y), EX(w), EY(h)); sp.shadow.inherit = False
    sp.line.fill.background()
    sp.fill.gradient()
    try:
        stops = sp.fill.gradient_stops
        stops[0].position = 0.0; stops[0].color.rgb = c1
        stops[1].position = 1.0; stops[1].color.rgb = c2
        sp.fill.gradient_angle = 45.0
    except Exception:
        sp.fill.solid(); sp.fill.fore_color.rgb = c1
    sp._element.spPr.append(parse_xml('<a:effectLst %s><a:softEdge rad="360000"/></a:effectLst>' % nsdecls('a')))
    return sp

def button(slide, x, y, w, h, label, fill, fg, size=12, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = rect(slide, x, y, w, h, fill=fill, shape=shape)
    tf = sp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = True
    r.font.color.rgb = fg; setfont(r)
    return sp

def pic(slide, path, x, y, w, h, soft=0):
    p = slide.shapes.add_picture(path, EX(x), EY(y), EX(w), EY(h))
    if soft > 0:
        p._element.spPr.append(parse_xml('<a:effectLst %s><a:softEdge rad="%d"/></a:effectLst>' % (nsdecls('a'), soft)))
    return p

def new_slide(page, code):
    s = prs.slides.add_slide(BLANK)
    bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    for r in range(3):
        for c in range(3):
            rect(s, 58 + c * 9, 44 + r * 9, 5, 5, fill=BLACK, shape=MSO_SHAPE.OVAL)
    text(s, 92, 38, 260, 64, [[("AIxlab", 17, True, BLACK)], [("BAIWUCHANG", 8, False, MUTED)]])
    text(s, 600, 44, 470, 52,
         [[("TEMPLATE      思源黑体      强调色", 10, True, BLACK)],
          [("通用模板       Bold + Light      #39FF14", 9, False, MUTED)]], align=PP_ALIGN.CENTER)
    text(s, 1380, 40, 170, 44, [[(page, 15, True, BLACK)]], align=PP_ALIGN.RIGHT)
    button(s, 1566, 40, 42, 42, "→", GREEN, BLACK, size=15)
    rect(s, 0, 115, 1672, 2, fill=LINE)
    rect(s, 0, 879, 1672, 2, fill=LINE)
    pic(s, LOGO, 56, 898, int(24 * 4.41), 24)
    text(s, 1430, 898, 186, 30, [[(code, 12, True, MUTED)]], align=PP_ALIGN.RIGHT)
    return s

def kicker(s, x, y, label):
    text(s, x, y, 600, 36, [[(label, 12, True, BLACK)]])
    rect(s, x, y + 34, 60, 6, fill=GREEN)

# ---------- 1 封面 ----------
s = new_slide("01 / 06", "T-01")
kicker(s, 56, 250, "通用模板 · TEMPLATE")
text(s, 54, 300, 820, 250, [[("在此输入", 52, True, BLACK)], [("演示主标题", 52, True, BLACK)]], line=1.05)
text(s, 56, 560, 660, 90, [[("副标题占位：用一句话清晰描述本次演示的主题与价值，保持简洁有力。", 14, False, DARK)]], line=1.45)
button(s, 56, 690, 168, 50, "开始 / START", GREEN, BLACK, size=13)
gblob(s, 1070, 250, 470, 470, GREEN, RGBColor(0xd6, 0xff, 0xc8))
pic(s, PLACE, 1095, 275, 420, 420, soft=180000)

# ---------- 2 章节分隔 ----------
s = new_slide("02 / 06", "T-02")
text(s, 56, 320, 470, 300, [[("01", 120, True, GREEN)]])
kicker(s, 640, 250, "章节 · SECTION")
text(s, 640, 300, 940, 110, [[("章节标题占位", 48, True, BLACK)]])
text(s, 640, 420, 900, 70, [[("本章节概述：用一两句话说明这一部分要讲的内容范围与重点。", 14, False, DARK)]], line=1.45)
rows = [("01", "要点条目一"), ("02", "要点条目二"), ("03", "要点条目三")]
yy = 545
for num, t in rows:
    rect(s, 640, yy + 18, 1660 - 640, 1, fill=LINE)
    rect(s, 648, yy + 18, 10, 10, fill=GREEN, shape=MSO_SHAPE.OVAL)
    text(s, 690, yy, 80, 40, [[(num, 16, True, BLACK)]])
    text(s, 760, yy, 500, 40, [[(t, 16, False, BLACK)]])
    yy += 70

# ---------- 3 大标题 + 正文 ----------
s = new_slide("03 / 06", "T-03")
kicker(s, 56, 200, "CHAPTER 01")
text(s, 54, 250, 700, 110, [[("大标题占位", 40, True, BLACK)]])
text(s, 56, 380, 720, 110, [[("正文段落占位：这里是常规内容区，用细体承载主要叙述。可写背景、方法或结论，建议每段控制在 2–4 行。", 14, False, DARK)]], line=1.6)
text(s, 56, 500, 720, 90, [[("第二段正文占位：右侧黑底高亮指标可与正文形成主次对比，让重点一眼可见。", 14, False, DARK)]], line=1.6)
rect(s, 1180, 150, 436, 600, fill=BLACK)
text(s, 1220, 360, 360, 120, [[("86%", 44, True, GREEN)]])
text(s, 1220, 470, 360, 40, [[("关键指标标题", 16, True, WHITE)]])
rect(s, 1220, 520, 356, 1, fill=RGBColor(0x55, 0x55, 0x55))
text(s, 1220, 540, 360, 120, [[("对该指标的简短说明：一句话解释它代表什么、为什么重要。", 12, False, RGBColor(0xc8, 0xc8, 0xc8))]], line=1.5)

# ---------- 4 三栏要点 ----------
s = new_slide("04 / 06", "T-04")
text(s, 54, 170, 700, 100, [[("三栏要点", 40, True, BLACK)]])
text(s, 1000, 190, 616, 80, [[("用三张卡片并列呈现要点，可将其中一张设为强调色高亮。", 13, False, DARK)]], line=1.5)
cards = [("01", "要点标题一", "要点说明占位：用一两句话描述该要点的核心内容。", False),
         ("02", "要点标题二", "高亮卡片：用强调色强调三项中最重要的一项。", True),
         ("03", "要点标题三", "要点说明占位：用一两句话描述该要点的核心内容。", False)]
cw = 500; gap = 30; x0 = 56; cy = 330; ch = 430
for i, (num, h2, body, active) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, cy, cw, ch, fill=(GREEN if active else WHITE), lc=(None if active else LINE), lw=1)
    text(s, x + 34, cy + 34, cw - 60, 36, [[(num, 16, True, (BLACK if active else MUTED))]])
    text(s, x + 34, cy + 78, cw - 60, 50, [[(h2, 19, True, BLACK)]])
    text(s, x + 34, cy + 140, cw - 60, 160, [[(body, 12.5, False, (BLACK if active else DARK))]], line=1.6)
    rect(s, x + 34, cy + ch - 80, 46, 46, fill=BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

# ---------- 5 关键数据 ----------
s = new_slide("05 / 06", "T-05")
kicker(s, 56, 170, "关键数据 · METRICS")
text(s, 54, 222, 700, 100, [[("用数据说话", 40, True, BLACK)]])
stats = [("86%", "指标一", 0.86, False), ("3.6x", "指标二", 0.72, True),
         ("127", "指标三", 0.64, False), ("99%", "指标四", 0.99, False)]
sw = 360; gap = 30; x0 = 56; sy = 380
for i, (num, label, pct, active) in enumerate(stats):
    x = x0 + i * (sw + gap)
    rect(s, x, sy, sw, 3, fill=BLACK)
    text(s, x, sy + 16, sw, 80, [[(num, 44, True, (GREEN if active else BLACK))]])
    text(s, x, sy + 110, sw, 34, [[(label, 15, True, BLACK)]])
    text(s, x, sy + 150, sw, 30, [[("简短说明文字", 11, False, MUTED)]])
    rect(s, x, sy + 200, sw, 8, fill=SOFTG)
    rect(s, x, sy + 200, int(sw * pct), 8, fill=GREEN)

# ---------- 6 结束页 ----------
s = new_slide("06 / 06", "T-06")
kicker(s, 56, 250, "结束 · THE END")
text(s, 54, 320, 1100, 160, [[("谢谢", 60, True, BLACK), ("观看", 60, True, GREEN)]])
text(s, 56, 520, 760, 70, [[("结束语占位：可在此致谢、留下一句号召或总结，并附上联系方式。", 14, False, DARK)]], line=1.45)
contact = [("EMAIL", "hello@aixlab.com"), ("WEB", "aixlab.com"), ("TEAM", "AIxlab · Baiwuchang")]
cx = 56
for b, v in contact:
    text(s, cx, 640, 320, 30, [[(b, 11, True, MUTED)]])
    text(s, cx, 670, 320, 40, [[(v, 16, False, BLACK)]])
    cx += 300

prs.save(OUT)
print("OK ->", OUT)
print("slides:", len(prs.slides._sldIdLst))
